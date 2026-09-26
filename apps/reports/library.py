"""The report library: what each report contains (rows from live data), and rendering to PDF or XLSX.

Every builder takes the page's filters (``Filters``) and returns a ``Table``. ``render`` turns a table into bytes.
"""

import io
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal

from django.db.models import Count, F, Q, Sum
from django.utils import timezone

from apps.academics.models import ClassGroup, Student, StudentGuardian


@dataclass
class Filters:
    start: date
    end: date
    grades: list = field(default_factory=list)
    sections: list = field(default_factory=list)  # ClassGroup ids
    term: str | None = None

    def groups(self):
        qs = ClassGroup.objects.all()
        if self.sections:
            qs = qs.filter(id__in=self.sections)
        elif self.grades:
            qs = qs.filter(grade__in=self.grades)
        return qs

    def narrowed(self) -> bool:
        return bool(self.sections or self.grades)

    def describe(self) -> str:
        scope = "All grades"
        if self.sections:
            scope = ", ".join(sorted(g.short_label for g in ClassGroup.objects.filter(id__in=self.sections)))
        elif self.grades:
            scope = "Grades " + ", ".join(self.grades)
        return f"{self.term + ' · ' if self.term else ''}{scope} · {self.start:%d %b} – {self.end:%d %b %Y}"

    def as_dict(self) -> dict:
        return {"start": self.start.isoformat(), "end": self.end.isoformat(), "grades": self.grades, "sections": [str(s) for s in self.sections], "term": self.term}

    @classmethod
    def from_dict(cls, d: dict) -> "Filters":
        return cls(start=date.fromisoformat(d["start"]), end=date.fromisoformat(d["end"]), grades=d.get("grades") or [], sections=d.get("sections") or [], term=d.get("term"))


@dataclass
class Table:
    title: str
    columns: list[str]
    rows: list[list]
    subtitle: str = ""
    widths: list[float] | None = None  # relative, for the PDF


def _pct(part, whole):
    return round(part * 100 / whole, 1) if whole else None


def _money(v) -> float:
    return float(v or 0)


# ------------------------------------------------------------------ builders


def attendance_register(f: Filters) -> Table:
    """Daily register by section: present, late, leave (excused), absent."""
    from apps.attendance.models import AttendanceException, AttendanceSession

    groups = {g.id: g for g in f.groups()}
    sizes = Counter(Student.objects.filter(is_active=True, class_group_id__in=groups).values_list("class_group_id", flat=True))
    sessions = AttendanceSession.objects.filter(class_group_id__in=groups, date__gte=f.start, date__lte=f.end).order_by("date")
    exc = defaultdict(Counter)
    for sid, st in AttendanceException.objects.filter(session__in=sessions).values_list("session_id", "status"):
        exc[sid][st] += 1
    rows = []
    for s in sessions.select_related("marked_by"):
        g = groups[s.class_group_id]
        total = sizes.get(g.id, 0)
        c = exc[s.id]
        absent, leave, late = c["absent"], c["excused"], c["late"] + c["half_day"]
        present = total - absent - leave
        rows.append([s.date.isoformat(), g.short_label, total, present - late, late, leave, absent, _pct(present, total), s.marked_by.full_name if s.marked_by else ""])
    rows.sort(key=lambda r: (r[0], _grade_order(r[1])))
    return Table("Attendance register", ["Date", "Section", "On roll", "Present", "Late (L)", "Leave (E)", "Absent (A)", "% present", "Marked by"], rows, widths=[1.1, 0.8, 0.8, 0.8, 0.8, 0.8, 0.8, 0.9, 1.6])


def _grade_order(label: str) -> tuple:
    from apps.principal.views import _grade_key

    grade, _, section = label.rpartition("-")
    return (_grade_key(grade), section)


def class_performance(f: Filters) -> Table:
    """Unit test and exam analysis by subject and section."""
    from apps.results.models import ExamMark

    marks = (
        ExamMark.objects.filter(exam__class_group__in=f.groups(), exam__is_published=True, exam__held_on__lte=f.end, is_absent=False)
        .values("exam__name", "exam__held_on", "exam__class_group__grade", "exam__class_group__section", "subject__name")
        .annotate(n=Count("id"), m=Sum("marks"), o=Sum("max_marks"), passed=Count("id", filter=Q(marks__gte=F("max_marks") * Decimal("0.33"))))
    )
    rows = []
    for r in marks:
        avg = _pct(float(r["m"]), float(r["o"]))
        rows.append([r["exam__name"], f"{r['exam__class_group__grade']}-{r['exam__class_group__section']}", r["subject__name"], r["n"], avg, _pct(r["passed"], r["n"])])
    rows.sort(key=lambda x: (x[0], _grade_order(x[1]), x[2]))
    return Table("Class performance", ["Exam", "Section", "Subject", "Students", "Average %", "Pass %"], rows, widths=[1.3, 0.8, 1.4, 0.8, 0.8, 0.8])


def fee_collection(f: Filters) -> Table:
    """Collections by mode, grade and date, with receipts."""
    from apps.fees.models import Payment

    pays = (
        Payment.objects.filter(status=Payment.Status.SUCCEEDED, paid_at__date__gte=f.start, paid_at__date__lte=f.end, invoice__student__class_group__in=f.groups())
        .select_related("invoice__student__class_group")
        .order_by("paid_at")
    )
    rows = [
        [timezone.localtime(p.paid_at).date().isoformat(), p.receipt_no or "", p.invoice.student.full_name, p.invoice.student.class_group.short_label, p.invoice.title, (p.method or p.gateway).upper(), _money(p.amount)]
        for p in pays
    ]
    total = sum(r[-1] for r in rows)
    rows.append(["", "", "", "", "Total", "", total])
    return Table("Fee collection", ["Date", "Receipt", "Student", "Class", "Invoice", "Mode", "Amount (₹)"], rows, widths=[0.9, 1.1, 1.5, 0.6, 1.6, 0.8, 0.9])


def defaulters(f: Filters) -> Table:
    """Overdue families with ageing and reminder history."""
    from apps.accounts.models import AuditLog
    from apps.fees.models import FeeInvoice
    from apps.notifications.models import Notification

    today = f.end
    invoices = (
        FeeInvoice.objects.filter(due_date__lt=today, student__class_group__in=f.groups(), student__is_active=True)
        .exclude(paid_amount__gte=F("amount"))
        .select_related("student__class_group")
        .order_by("due_date")
    )
    guardians = defaultdict(list)
    for link in StudentGuardian.objects.filter(student__in=invoices.values("student")).select_related("user"):
        guardians[link.student_id].append(link.user)
    rows = []
    for inv in invoices:
        days = (today - inv.due_date).days
        bucket = "0–30 days" if days <= 30 else "31–60 days" if days <= 60 else "61–90 days" if days <= 90 else "90+ days"
        users = guardians.get(inv.student_id, [])
        reminders = Notification.objects.filter(user__in=users, category="fees").count() + AuditLog.objects.filter(action__startswith="fees.remind", target_id=str(inv.id)).count()
        rows.append([inv.student.full_name, inv.student.class_group.short_label, users[0].full_name if users else "", inv.title, inv.due_date.isoformat(), days, bucket, _money(inv.amount - inv.paid_amount), reminders])
    return Table("Defaulters", ["Student", "Class", "Guardian", "Invoice", "Due", "Days overdue", "Ageing", "Balance (₹)", "Reminders"], rows)


def staff_attendance(f: Filters) -> Table:
    """Monthly muster per staff member: present, late, absent, leave; leave taken and substitutions covered."""
    from apps.accounts.models import Membership, Role
    from apps.staff.models import StaffAttendance, StaffLeave, Substitution

    rows = []
    marks = defaultdict(Counter)
    for uid, st in StaffAttendance.objects.filter(date__gte=f.start, date__lte=f.end).values_list("user_id", "status"):
        marks[uid][st] += 1
    leave_days = defaultdict(Decimal)
    for uid, days in StaffLeave.objects.filter(status="approved", from_date__lte=f.end, to_date__gte=f.start).values_list("user_id", "days"):
        leave_days[uid] += days
    covers = Counter(Substitution.objects.filter(date__gte=f.start, date__lte=f.end).values_list("teacher_id", flat=True))
    seen = set()
    for m in Membership.objects.filter(is_active=True).exclude(role__in=[Role.PARENT, Role.STUDENT]).select_related("user").order_by("user__full_name"):
        if m.user_id in seen:
            continue
        seen.add(m.user_id)
        c = marks.get(m.user_id, Counter())
        total = sum(c.values())
        if not total:
            continue
        rows.append([m.user.full_name, m.title or m.get_role_display(), c["present"], c["late"], c["absent"], c["leave"], _pct(c["present"] + c["late"], total), float(leave_days.get(m.user_id, 0)), covers.get(m.user_id, 0)])
    return Table("Staff attendance", ["Name", "Role", "Present", "Late", "Absent", "Leave", "% present", "Leave days", "Covers"], rows, widths=[1.6, 1.6, 0.7, 0.6, 0.7, 0.6, 0.8, 0.8, 0.6])


def transport_utilisation(f: Filters) -> Table:
    """Seat use and punctuality by route and run."""
    from datetime import datetime

    from apps.transport.models import Route, StudentTransport, Trip

    riders = Counter(StudentTransport.objects.filter(is_active=True, student__class_group__in=f.groups()).values_list("route_id", flat=True))
    rows = []
    for r in Route.objects.filter(is_active=True).select_related("vehicle").order_by("code"):
        cap = r.vehicle.capacity if r.vehicle else 0
        trips = Trip.objects.filter(route=r, service_date__gte=f.start, service_date__lte=f.end)
        started = [t for t in trips if t.started_at]
        delays = []
        for t in started:
            planned = datetime.combine(t.service_date, t.scheduled_start, tzinfo=timezone.localtime(t.started_at).tzinfo)
            delays.append(max(0, (timezone.localtime(t.started_at) - planned).total_seconds() / 60))
        on_time = sum(1 for d in delays if d <= 5)
        rows.append([r.name, r.vehicle.registration_no if r.vehicle else "", cap, riders.get(r.id, 0), _pct(riders.get(r.id, 0), cap), trips.count(), len(started), _pct(on_time, len(delays)), round(sum(delays) / len(delays), 1) if delays else None])
    return Table("Transport utilisation", ["Route", "Vehicle", "Seats", "Riders", "Seat use %", "Runs planned", "Runs made", "On time %", "Avg delay (min)"], rows)


def admissions_funnel(f: Filters) -> Table:
    """Enquiry → application → admitted, for the admissions cycle, by grade and source."""
    from apps.admissions.models import AdmissionCycle, Application

    cycle = AdmissionCycle.objects.filter(is_current=True).first()
    apps = Application.objects.all()
    if cycle:
        apps = apps.filter(academic_year=cycle.academic_year)
    if f.grades:
        apps = apps.filter(grade__in=f.grades)
    stages = [s.value for s in Application.Stage]
    reached = {s: i for i, s in enumerate(stages)}
    rows = []
    by = defaultdict(Counter)
    for a in apps.values("grade", "source", "stage", "closed"):
        key = (a["grade"], Application.Source(a["source"]).label if a["source"] in Application.Source.values else (a["source"] or "—"))
        at = reached.get(a["stage"], 0)
        c = by[key]
        c["enquiries"] += 1
        c["applications"] += at >= reached["application"]
        c["assessed"] += at >= reached["assessment"]
        c["offers"] += at >= reached["offer"]
        c["admitted"] += at >= reached["admitted"]
        c["closed"] += bool(a["closed"])
    from apps.principal.views import _grade_key

    for (grade, source), c in sorted(by.items(), key=lambda kv: (_grade_key(kv[0][0]), kv[0][1])):
        rows.append([grade, source, c["enquiries"], c["applications"], c["assessed"], c["offers"], c["admitted"], _pct(c["admitted"], c["enquiries"]), c["closed"]])
    tot = Counter()
    for c in by.values():
        tot.update(c)
    rows.append(["All", "", tot["enquiries"], tot["applications"], tot["assessed"], tot["offers"], tot["admitted"], _pct(tot["admitted"], tot["enquiries"]), tot["closed"]])
    title = f"Admissions funnel {cycle.academic_year}" if cycle else "Admissions funnel"
    return Table(title, ["Grade", "Source", "Enquiries", "Applications", "Assessed", "Offers", "Admitted", "Conversion %", "Closed"], rows)


LIBRARY = {
    "attendance_register": {"build": attendance_register, "formats": ["pdf", "xlsx"], "module": "attendance"},
    "class_performance": {"build": class_performance, "formats": ["pdf", "xlsx"], "module": "exams"},
    "fee_collection": {"build": fee_collection, "formats": ["pdf", "xlsx"], "module": "fees"},
    "defaulters": {"build": defaulters, "formats": ["xlsx"], "module": "fees"},
    "staff_attendance": {"build": staff_attendance, "formats": ["pdf", "xlsx"], "module": "staff"},
    "transport_utilisation": {"build": transport_utilisation, "formats": ["pdf"], "module": "transport"},
    "admissions_funnel": {"build": admissions_funnel, "formats": ["pdf", "xlsx"], "module": "admissions"},
}


# ------------------------------------------------------------------ custom reports

CUSTOM_MODULES = {
    "students": {
        "admission_no": "Admission no.",
        "name": "Name",
        "class": "Class",
        "roll_no": "Roll no.",
        "gender": "Gender",
        "date_of_birth": "Date of birth",
        "house": "House",
        "guardian": "Guardian",
        "route": "Bus route",
        "attendance": "Attendance % (range)",
    },
    "staff": {"name": "Name", "role": "Role", "title": "Title", "department": "Department", "attendance": "Attendance % (range)", "leave_days": "Leave days (range)"},
    "fees": {"student": "Student", "class": "Class", "invoice": "Invoice", "category": "Category", "amount": "Amount (₹)", "paid": "Paid (₹)", "balance": "Balance (₹)", "due_date": "Due date"},
    "attendance": {"date": "Date", "section": "Section", "on_roll": "On roll", "present": "Present", "late": "Late", "absent": "Absent", "percent": "% present"},
}


def custom(module: str, columns: list[str], f: Filters, name: str) -> Table:
    catalog = CUSTOM_MODULES[module]
    cols = [c for c in columns if c in catalog]
    if module == "students":
        from apps.transport.models import StudentTransport

        students = list(Student.objects.filter(is_active=True, class_group__in=f.groups()).select_related("class_group").order_by("class_group__grade", "class_group__section", "roll_no"))
        guardian = {}
        for link in StudentGuardian.objects.filter(student__in=students).select_related("user").order_by("-is_primary"):
            guardian.setdefault(link.student_id, link.user.full_name)
        route = dict(StudentTransport.objects.filter(student__in=students, is_active=True).values_list("student_id", "route__name"))
        att = _student_attendance(students, f) if "attendance" in cols else {}
        values = lambda s: {  # noqa: E731
            "admission_no": s.admission_no, "name": s.full_name, "class": s.class_group.short_label, "roll_no": s.roll_no, "gender": s.gender,
            "date_of_birth": s.date_of_birth.isoformat() if s.date_of_birth else "", "house": s.house, "guardian": guardian.get(s.id, ""),
            "route": route.get(s.id, ""), "attendance": att.get(s.id),
        }
        rows = [[values(s)[c] for c in cols] for s in students]
    elif module == "staff":
        table = staff_attendance(f)
        by_name = {r[0]: r for r in table.rows}
        from apps.accounts.models import Membership, Role

        rows, seen = [], set()
        for m in Membership.objects.filter(is_active=True).exclude(role__in=[Role.PARENT, Role.STUDENT]).select_related("user").order_by("user__full_name"):
            if m.user_id in seen:
                continue
            seen.add(m.user_id)
            r = by_name.get(m.user.full_name)
            v = {"name": m.user.full_name, "role": m.get_role_display(), "title": m.title, "department": m.get_department_display() if m.department else "", "attendance": r[6] if r else None, "leave_days": r[7] if r else 0}
            rows.append([v[c] for c in cols])
    elif module == "fees":
        from apps.fees.models import FeeInvoice

        rows = []
        for inv in FeeInvoice.objects.filter(student__class_group__in=f.groups(), due_date__gte=f.start, due_date__lte=f.end).select_related("student__class_group").order_by("due_date"):
            v = {"student": inv.student.full_name, "class": inv.student.class_group.short_label, "invoice": inv.title, "category": inv.get_category_display(), "amount": _money(inv.amount), "paid": _money(inv.paid_amount), "balance": _money(inv.amount - inv.paid_amount), "due_date": inv.due_date.isoformat()}
            rows.append([v[c] for c in cols])
    else:
        reg = attendance_register(f)
        keys = ["date", "section", "on_roll", "present", "late", "leave", "absent", "percent"]
        rows = []
        for r in reg.rows:
            v = dict(zip(keys, r))
            v["present"] = r[3] + r[4]
            rows.append([v[c] for c in cols])
    return Table(name, [catalog[c] for c in cols], rows)


def _student_attendance(students, f: Filters) -> dict:
    from apps.attendance.models import AttendanceException, AttendanceSession

    days = Counter(AttendanceSession.objects.filter(class_group__in={s.class_group_id for s in students}, date__gte=f.start, date__lte=f.end).values_list("class_group_id", flat=True))
    off = Counter(
        AttendanceException.objects.filter(session__date__gte=f.start, session__date__lte=f.end, status__in=["absent", "excused"], student__in=students).values_list("student_id", flat=True)
    )
    return {s.id: _pct(days[s.class_group_id] - off[s.id], days[s.class_group_id]) for s in students}


# ------------------------------------------------------------------ rendering


def render(table: Table, fmt: str, school_name: str, filters: Filters) -> bytes:
    return _xlsx(table, school_name, filters) if fmt == "xlsx" else _pdf(table, school_name, filters)


def _xlsx(table: Table, school_name: str, filters: Filters) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = table.title[:30]
    ws.append([f"{school_name} · {table.title}"])
    ws.append([filters.describe()])
    ws.append([f"Generated {timezone.localtime():%d %b %Y, %I:%M %p}"])
    ws.append([])
    ws.append(table.columns)
    head = ws.max_row
    for cell in ws[head]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3446C8")
        cell.alignment = Alignment(vertical="center")
    ws["A1"].font = Font(bold=True, size=14)
    for row in table.rows:
        ws.append(row)
    for i, name in enumerate(table.columns, start=1):
        width = max([len(str(name))] + [len(str(r[i - 1])) for r in table.rows[:500] if i - 1 < len(r)])
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(10, width + 2))
    ws.freeze_panes = ws.cell(row=head + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf(table: Table, school_name: str, filters: Filters) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    buf = io.BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm, title=table.title)
    styles = getSampleStyleSheet()
    story = [
        Paragraph(f"<b>{_esc(school_name)}</b> · {_esc(table.title)}", styles["Title"]),
        Paragraph(_esc(filters.describe()), styles["Normal"]),
        Paragraph(f"Generated {timezone.localtime():%d %b %Y, %I:%M %p} · {len(table.rows)} rows", styles["Normal"]),
        Spacer(1, 6 * mm),
    ]
    width = page[0] - 28 * mm
    rel = table.widths or [1] * len(table.columns)
    col_widths = [width * w / sum(rel) for w in rel]
    cell = styles["BodyText"].clone("cell", fontSize=8, leading=10)
    data = [[Paragraph(f"<b>{_esc(c)}</b>", cell) for c in table.columns]] + [[Paragraph(_esc(_fmt(v)), cell) for v in r] for r in table.rows] if table.rows else [[Paragraph(_esc(c), cell) for c in table.columns], [Paragraph("No rows for these filters.", cell)] + [""] * (len(table.columns) - 1)]
    t = LongTable(data, colWidths=col_widths, repeatRows=1)
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EAF8")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F4EF")]),
                ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor("#3446C8")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story.append(t)
    doc.build(story)
    return buf.getvalue()


def _fmt(v) -> str:
    if v is None:
        return "—"
    if isinstance(v, float):
        return f"{v:,.1f}" if not v.is_integer() else f"{v:,.0f}"
    return str(v)


def _esc(s) -> str:
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

